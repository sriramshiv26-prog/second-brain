"""Multi-format document parser for the Second Brain system."""

import logging
import re
from pathlib import Path
from typing import Dict, Any, Optional
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup

# Configure logging
logger = logging.getLogger(__name__)

# Constants
MAX_TEXT_LENGTH = 100000  # 100,000 character limit


class DocumentParser:
    """
    Parser for multiple document formats including text, markdown, PDF, DOCX, XLSX, and URLs.

    Provides a unified interface for extracting text and metadata from various sources.
    """

    def __init__(self) -> None:
        """Initialize the DocumentParser."""
        logger.info("Initializing DocumentParser")

    # ========================================================================
    # Public API: Main parse method (dispatcher)
    # ========================================================================

    def parse(self, source: str) -> Dict[str, Any]:
        """
        Parse a document from a file path or URL.

        Automatically detects if the source is a URL or file path and calls the
        appropriate parser.

        Args:
            source: File path (str or Path) or URL (str)

        Returns:
            Dictionary with keys: text, title, source, format, and optionally author, date

        Example:
            >>> parser = DocumentParser()
            >>> result = parser.parse("https://example.com")
            >>> result = parser.parse("/path/to/document.pdf")
        """
        try:
            # Check if source is a URL
            if self._is_url(source):
                logger.info(f"Parsing URL: {source}")
                return self.parse_url(source)

            # Otherwise, treat as file path
            file_path = Path(source)

            if not file_path.exists():
                logger.error(f"File not found: {file_path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": str(file_path),
                    "format": "unknown",
                    "error": f"File not found: {file_path}",
                }

            # Get file extension
            suffix = file_path.suffix.lower()
            logger.info(f"Parsing file: {file_path} (format: {suffix})")

            # Dispatch to appropriate parser
            if suffix == ".txt":
                return self.parse_text(str(file_path))
            elif suffix == ".md":
                return self.parse_markdown(str(file_path))
            elif suffix == ".pdf":
                return self.parse_pdf(str(file_path))
            elif suffix == ".docx":
                return self.parse_docx(str(file_path))
            elif suffix == ".xlsx":
                return self.parse_xlsx(str(file_path))
            else:
                logger.warning(f"Unsupported file format: {suffix}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": str(file_path),
                    "format": "unknown",
                    "error": f"Unsupported file format: {suffix}",
                }

        except Exception as e:
            logger.error(f"Error in parse dispatcher: {e}", exc_info=True)
            return {
                "text": "",
                "title": str(source),
                "source": str(source),
                "format": "unknown",
                "error": str(e),
            }

    # ========================================================================
    # Format-specific parsers
    # ========================================================================

    def parse_text(self, path: str) -> Dict[str, Any]:
        """
        Parse a plain text file.

        Args:
            path: File path to the text file

        Returns:
            Dictionary with text content and metadata
        """
        try:
            file_path = Path(path)

            if not file_path.exists():
                logger.error(f"Text file not found: {path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": path,
                    "format": "text",
                    "error": f"File not found: {path}",
                }

            # Read file with error handling for encoding issues
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

            # Limit text length
            text = self._limit_text(text)

            logger.info(f"Successfully parsed text file: {path} ({len(text)} chars)")

            return {
                "text": text,
                "title": file_path.stem,
                "source": path,
                "format": "text",
                "author": None,
                "date": None,
            }

        except Exception as e:
            logger.error(f"Error parsing text file {path}: {e}", exc_info=True)
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "text",
                "error": str(e),
            }

    def parse_markdown(self, path: str) -> Dict[str, Any]:
        """
        Parse a markdown file.

        Args:
            path: File path to the markdown file

        Returns:
            Dictionary with text content and metadata
        """
        try:
            file_path = Path(path)

            if not file_path.exists():
                logger.error(f"Markdown file not found: {path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": path,
                    "format": "markdown",
                    "error": f"File not found: {path}",
                }

            # Read file with error handling for encoding issues
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

            # Extract title from first heading if present
            title = self._extract_markdown_title(text, file_path.stem)

            # Limit text length
            text = self._limit_text(text)

            logger.info(f"Successfully parsed markdown file: {path} ({len(text)} chars)")

            return {
                "text": text,
                "title": title,
                "source": path,
                "format": "markdown",
                "author": None,
                "date": None,
            }

        except Exception as e:
            logger.error(f"Error parsing markdown file {path}: {e}", exc_info=True)
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "markdown",
                "error": str(e),
            }

    def parse_pdf(self, path: str) -> Dict[str, Any]:
        """
        Parse a PDF file.

        Args:
            path: File path to the PDF file

        Returns:
            Dictionary with text content and metadata
        """
        try:
            import pymupdf

            file_path = Path(path)

            if not file_path.exists():
                logger.error(f"PDF file not found: {path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": path,
                    "format": "pdf",
                    "error": f"File not found: {path}",
                }

            # Open and extract text from PDF
            doc = pymupdf.open(path)
            text = ""

            for page in doc:
                text += page.get_text()

            doc.close()

            # Extract metadata
            metadata = doc.metadata
            author = metadata.get("author") if metadata else None
            title = metadata.get("title") if metadata else file_path.stem

            # Limit text length
            text = self._limit_text(text)

            logger.info(f"Successfully parsed PDF file: {path} ({len(text)} chars)")

            return {
                "text": text,
                "title": title or file_path.stem,
                "source": path,
                "format": "pdf",
                "author": author,
                "date": None,
            }

        except ImportError:
            logger.error("pymupdf not installed. Install with: pip install pymupdf")
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "pdf",
                "error": "pymupdf not installed",
            }
        except Exception as e:
            logger.error(f"Error parsing PDF file {path}: {e}", exc_info=True)
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "pdf",
                "error": str(e),
            }

    def parse_docx(self, path: str) -> Dict[str, Any]:
        """
        Parse a DOCX (Microsoft Word) file.

        Args:
            path: File path to the DOCX file

        Returns:
            Dictionary with text content and metadata
        """
        try:
            from docx import Document

            file_path = Path(path)

            if not file_path.exists():
                logger.error(f"DOCX file not found: {path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": path,
                    "format": "docx",
                    "error": f"File not found: {path}",
                }

            # Open and extract text from DOCX
            doc = Document(path)
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

            # Extract metadata
            properties = doc.core_properties
            author = properties.author
            title = properties.title or file_path.stem

            # Limit text length
            text = self._limit_text(text)

            logger.info(f"Successfully parsed DOCX file: {path} ({len(text)} chars)")

            return {
                "text": text,
                "title": title,
                "source": path,
                "format": "docx",
                "author": author,
                "date": None,
            }

        except ImportError:
            logger.error("python-docx not installed. Install with: pip install python-docx")
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "docx",
                "error": "python-docx not installed",
            }
        except Exception as e:
            logger.error(f"Error parsing DOCX file {path}: {e}", exc_info=True)
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "docx",
                "error": str(e),
            }

    def parse_xlsx(self, path: str) -> Dict[str, Any]:
        """
        Parse an XLSX (Microsoft Excel) file.

        Args:
            path: File path to the XLSX file

        Returns:
            Dictionary with text content and metadata
        """
        try:
            from openpyxl import load_workbook

            file_path = Path(path)

            if not file_path.exists():
                logger.error(f"XLSX file not found: {path}")
                return {
                    "text": "",
                    "title": file_path.name,
                    "source": path,
                    "format": "xlsx",
                    "error": f"File not found: {path}",
                }

            # Open and extract text from XLSX
            wb = load_workbook(path)
            text_parts = []

            for sheet in wb.sheetnames:
                ws = wb[sheet]
                text_parts.append(f"Sheet: {sheet}\n")
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    text_parts.append(row_text)
                text_parts.append("\n")

            text = "\n".join(text_parts)

            # Limit text length
            text = self._limit_text(text)

            logger.info(f"Successfully parsed XLSX file: {path} ({len(text)} chars)")

            return {
                "text": text,
                "title": file_path.stem,
                "source": path,
                "format": "xlsx",
                "author": None,
                "date": None,
            }

        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "xlsx",
                "error": "openpyxl not installed",
            }
        except Exception as e:
            logger.error(f"Error parsing XLSX file {path}: {e}", exc_info=True)
            return {
                "text": "",
                "title": Path(path).name,
                "source": path,
                "format": "xlsx",
                "error": str(e),
            }

    def parse_url(self, url: str) -> Dict[str, Any]:
        """
        Parse a URL and extract text content.

        Fetches the HTML from the URL, removes script/style/nav/footer tags,
        and extracts readable text content.

        Args:
            url: The URL to parse

        Returns:
            Dictionary with text content and metadata
        """
        try:
            logger.info(f"Fetching URL: {url}")

            # Fetch the page
            headers = {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
                )
            }
            response = requests.get(url, headers=headers, timeout=10)
            response.raise_for_status()

            # Parse HTML
            soup = BeautifulSoup(response.text, "html.parser")

            # Remove script, style, nav, footer tags
            for tag in soup.find_all(["script", "style", "nav", "footer"]):
                tag.decompose()

            # Extract text
            text = soup.get_text(separator="\n", strip=True)

            # Limit text length
            text = self._limit_text(text)

            # Extract title
            title = soup.title.string if soup.title else url

            logger.info(f"Successfully parsed URL: {url} ({len(text)} chars)")

            return {
                "text": text,
                "title": title,
                "source": url,
                "format": "html",
                "author": None,
                "date": None,
            }

        except requests.RequestException as e:
            logger.error(f"Error fetching URL {url}: {e}", exc_info=True)
            return {
                "text": "",
                "title": url,
                "source": url,
                "format": "html",
                "error": f"Failed to fetch URL: {str(e)}",
            }
        except Exception as e:
            logger.error(f"Error parsing URL {url}: {e}", exc_info=True)
            return {
                "text": "",
                "title": url,
                "source": url,
                "format": "html",
                "error": str(e),
            }

    # ========================================================================
    # Helper methods
    # ========================================================================

    @staticmethod
    def _is_url(source: str) -> bool:
        """
        Check if a source string is a URL.

        Args:
            source: The source string to check

        Returns:
            True if source is a URL, False otherwise
        """
        try:
            result = urlparse(source)
            return result.scheme in ("http", "https")
        except Exception:
            return False

    @staticmethod
    def _limit_text(text: str, max_length: int = MAX_TEXT_LENGTH) -> str:
        """
        Limit text to maximum length.

        Args:
            text: The text to limit
            max_length: Maximum length in characters

        Returns:
            Text limited to max_length characters
        """
        if len(text) > max_length:
            logger.warning(
                f"Text exceeds max length ({len(text)} > {max_length}), truncating"
            )
            return text[:max_length]
        return text

    @staticmethod
    def _extract_markdown_title(text: str, fallback: str = "") -> str:
        """
        Extract title from markdown text (first H1 heading).

        Args:
            text: The markdown text
            fallback: Fallback title if no heading found

        Returns:
            The extracted title or fallback
        """
        # Look for H1 heading
        match = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
        if match:
            return match.group(1).strip()
        return fallback
